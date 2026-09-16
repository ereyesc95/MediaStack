"""Excel import/export for the per-profile music collection."""
from __future__ import annotations

import io
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.music_collection import (
    _band_country,
    _band_genres,
    _find_band_by_name,
    find_existing,
    fuzzy_matches,
    match_key,
    normalize_media_type,
    parse_album_line,
    split_media_types,
    split_sources,
    upsert_item,
)

HEADERS = (
    "ALBUM",
    "RELEASE",
    "ANIMATION",
    "CANVAS",
    "MEDIA TYPE",
    "AUTOGRAPHS",
    "PENDING",
)


def _autograph_list(raw: str | None) -> list[str]:
    if not raw:
        return []
    parts = []
    for line in str(raw).replace("\r", "").split("\n"):
        for piece in re_split_autographs(line):
            if piece:
                parts.append(piece)
    # dedupe
    out: list[str] = []
    seen: set[str] = set()
    for p in parts:
        key = p.casefold()
        if key in seen:
            continue
        seen.add(key)
        out.append(p)
    return out


def re_split_autographs(line: str) -> list[str]:
    import re

    return [p.strip() for p in re.split(r"[/,|;]+", line) if p.strip()]


def _pending_list(raw: str | None) -> list[str]:
    if not raw:
        return []
    text = str(raw).strip()
    if not text:
        return []
    if text.casefold() == "everything":
        return ["Everything"]
    if text.casefold() in ("spotify code", "spotify - code", "spotify"):
        return ["Spotify Code"]
    return [p.strip() for p in text.replace(";", ",").split(",") if p.strip()]


def _album_export_line(item: dict) -> str:
    artist = item.get("artist") or ""
    title = item.get("title") or ""
    date = item.get("original_date") or item.get("year") or ""
    edition = item.get("edition") or "Standard Edition"
    base = f"{artist}, {date}. {title}" if date else f"{artist}, {title}"
    if edition and edition.casefold() != "standard edition":
        return f"{base}, {edition}"
    return base


def export_collection_xlsx(db: Session, user_id: int) -> bytes:
    from app.music_collection import list_items

    items = list_items(db, user_id, sort="artist", order="asc")
    wb = Workbook()
    ws = wb.active
    ws.title = "Main"
    ws.append(list(HEADERS))
    for it in items:
        anim = "/".join(it.get("animation") or []) or None
        canv = "/".join(it.get("canvas") or []) or None
        autos = "/".join(it.get("autographs") or []) or None
        pending = it.get("pending") or []
        pending_cell = None
        if pending:
            if len(pending) >= 10:
                pending_cell = "Everything"
            elif len(pending) == 1 and pending[0] == "Spotify - Card":
                pending_cell = "Spotify Code"
            else:
                pending_cell = ", ".join(pending)
        ws.append(
            [
                _album_export_line(it),
                it.get("release_type"),
                anim,
                canv,
                it.get("media_type"),
                autos,
                pending_cell,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview_import(db: Session, user_id: int, file_obj: BinaryIO) -> dict:
    wb = load_workbook(file_obj, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {"ok": False, "error": "Empty workbook", "rows": []}
    header = [str(c or "").strip().upper() for c in rows[0]]
    # Map columns
    idx = {name: header.index(name) if name in header else -1 for name in HEADERS}
    if idx["ALBUM"] < 0:
        return {"ok": False, "error": "Missing ALBUM column", "rows": []}

    preview_rows: list[dict] = []
    counts = {"add": 0, "update": 0, "orphan": 0, "skip": 0}

    for raw in rows[1:]:
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue

        def cell(name: str):
            i = idx[name]
            if i < 0 or i >= len(raw):
                return None
            v = raw[i]
            return None if v is None else str(v).strip()

        album = cell("ALBUM") or ""
        parsed = parse_album_line(album)
        release_type = cell("RELEASE") or "Studio Album"
        animation = split_sources(cell("ANIMATION"))
        canvas = split_sources(cell("CANVAS"))
        media_list = split_media_types(cell("MEDIA TYPE")) or [None]
        autographs = _autograph_list(cell("AUTOGRAPHS"))
        pending_extra = _pending_list(cell("PENDING"))

        band = _find_band_by_name(db, parsed["artist"])
        country, iso = _band_country(db, band) if band else (None, None)
        genres = _band_genres(db, band) if band else []
        matches = fuzzy_matches(db, artist=parsed["artist"], title=parsed["title"])

        for media in media_list:
            key = match_key(
                artist=parsed["artist"],
                title=parsed["title"],
                edition=parsed["edition"],
                media_type=media,
                date_iso=parsed["date_iso"],
                release_type=release_type,
            )
            existing = find_existing(db, user_id, key)
            action = "update" if existing else "add"
            orphan = band is None or not matches
            if orphan:
                counts["orphan"] += 1
            counts[action] += 1
            preview_rows.append(
                {
                    "action": action,
                    "orphan": orphan,
                    "artist": parsed["artist"],
                    "title": parsed["title"],
                    "edition": parsed["edition"],
                    "date_iso": parsed["date_iso"],
                    "release_type": release_type,
                    "media_type": normalize_media_type(media) if media else None,
                    "animation": animation,
                    "canvas": canvas,
                    "autographs": autographs,
                    "pending_extra": pending_extra,
                    "genres": genres,
                    "country": country,
                    "country_iso": iso,
                    "band_id": band.bnd_id if band else None,
                    "match_key": key,
                    "possible_matches": matches[:5],
                    "existing_id": existing.col_id if existing else None,
                }
            )

    return {
        "ok": True,
        "counts": counts,
        "rows": preview_rows,
        "total": len(preview_rows),
    }


def commit_import(db: Session, user_id: int, rows: list[dict]) -> dict:
    added = updated = 0
    for row in rows:
        payload = {
            "artist": row.get("artist"),
            "title": row.get("title"),
            "edition": row.get("edition") or "Standard Edition",
            "release_type": row.get("release_type"),
            "original_date": row.get("date_iso"),
            "media_type": row.get("media_type"),
            "animation": row.get("animation") or [],
            "canvas": row.get("canvas") or [],
            "autographs": row.get("autographs") or [],
            "pending_extra": row.get("pending_extra") or [],
            "genres": row.get("genres") or [],
            "country": row.get("country"),
            "country_iso": row.get("country_iso"),
            "band_id": row.get("band_id"),
            "match_key": row.get("match_key"),
        }
        # Link first possible match folder if present
        matches = row.get("possible_matches") or []
        if matches and not payload.get("folder_path"):
            m0 = matches[0]
            payload["folder_path"] = m0.get("folder_path")
            payload["release_folder_path"] = m0.get("folder_path")
            payload["release_id"] = m0.get("release_id")
            if m0.get("band_id"):
                payload["band_id"] = m0["band_id"]

        existing = find_existing(db, user_id, payload["match_key"])
        upsert_item(db, user_id, payload)
        if existing:
            updated += 1
        else:
            added += 1
    return {"ok": True, "added": added, "updated": updated, "total": added + updated}
